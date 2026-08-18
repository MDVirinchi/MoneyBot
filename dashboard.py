"""
MoneyBot RS60/EP40 - Operator Dashboard
Reads: ops_log_*.json + MoneyBot_PaperTrading_Package.xlsx
Output: dashboard.html opened in browser. Auto-refreshes every 5 min.

Run after daily_ops_report.py: python dashboard.py
"""

import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

import json, glob, os, webbrowser, csv
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl not found. Run: pip install openpyxl")
    sys.exit(1)

BASE          = os.path.dirname(os.path.abspath(__file__))
WORKBOOK      = os.path.join(BASE, "MoneyBot_PaperTrading_Package.xlsx")
try:
    import config as _cfg
    PAPER_CAPITAL = getattr(_cfg, "TRADING_CAPITAL_INR", 5_000)
except Exception:
    PAPER_CAPITAL = 5_000
PAPER_START   = date(2026, 6, 16)

# ── 1. OPS LOGS ──────────────────────────────────────────────────────────────
def load_ops_logs():
    files = sorted(glob.glob(os.path.join(BASE, "ops_log_*.json")))
    logs = []
    for f in files:
        try:
            with open(f) as fh:
                logs.append(json.load(fh))
        except Exception:
            pass
    return logs

# ── 2. WORKBOOK ───────────────────────────────────────────────────────────────
def load_workbook():
    result = {
        "positions": [], "trades": [],
        "rebalance_dates": [], "daily_log": [],
    }
    if not os.path.exists(WORKBOOK):
        return result
    try:
        wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    except Exception as e:
        print(f"WARNING: workbook error: {e}")
        return result

    # Position_Log: row 3 = header, rows 4+ = data then summary footer.
    # Real position rows have a datetime in col B (Entry Date).
    if "Position_Log" in wb.sheetnames:
        ws = wb["Position_Log"]
        for row in ws.iter_rows(min_row=4, values_only=True):
            if not isinstance(row[1], datetime):
                continue
            result["positions"].append({
                "symbol":    str(row[0]) if row[0] else "?",
                "entry_date": row[1].date() if row[1] else None,
                "entry_px":  row[2],
                "qty":       row[3],
                "sl":        row[4],
                "current":   row[5],
                "unrealized": row[6],
                "pnl_pct":   row[7],
                "days_held": row[8],
                "alloc":     row[9],
            })

    # Trade_Log: row 3 = header, rows 4+
    if "Trade_Log" in wb.sheetnames:
        ws = wb["Trade_Log"]
        for row in ws.iter_rows(min_row=4, values_only=True):
            if row[1] is None:
                continue
            result["trades"].append({
                "symbol":  row[1],
                "net_pnl": row[9],
                "win":     row[13],
            })

    # Rebalance_Log: col B (index 1) has the rebalance dates
    if "Rebalance_Log" in wb.sheetnames:
        ws = wb["Rebalance_Log"]
        for row in ws.iter_rows(min_row=4, values_only=True):
            d = row[1]
            if isinstance(d, datetime):
                result["rebalance_dates"].append(d.date())

    # Daily_Ops_Log: col B=date, col H=portfolio value (index 7)
    if "Daily_Ops_Log" in wb.sheetnames:
        ws = wb["Daily_Ops_Log"]
        for row in ws.iter_rows(min_row=4, values_only=True):
            d = row[1]
            if not isinstance(d, datetime):
                continue
            pv = row[7]
            result["daily_log"].append({
                "date":      d.date(),
                "portfolio": pv if isinstance(pv, (int, float)) else None,
            })

    return result

# ── 2b. SHADOW DATA ───────────────────────────────────────────────────────────
def load_shadow_data():
    result = {"regime": "UNKNOWN", "positions": 0, "pnl": 0.0,
              "portfolio_value": PAPER_CAPITAL, "return_pct": 0.0,
              "trades": 0, "wins": 0, "win_rate": 0.0,
              "comparison_verdict": "", "comparison_delta": 0.0}
    state_file = Path(BASE) / "shadow_state.json"
    if state_file.exists():
        try:
            state = json.loads(state_file.read_text(encoding="utf-8"))
            positions = state.get("positions", {})
            invested = sum(p.get("entry", 0) * p.get("qty", 0) for p in positions.values())
            pv = state.get("cash", float(PAPER_CAPITAL)) + invested
            result["positions"]       = len(positions)
            result["pnl"]             = round(state.get("pnl", 0), 2)
            result["portfolio_value"] = round(pv, 2)
            result["return_pct"]      = round((pv - PAPER_CAPITAL) / PAPER_CAPITAL * 100, 2)
        except Exception:
            pass
    journal_file = Path(BASE) / "shadow_journal.csv"
    if journal_file.exists():
        try:
            rows = list(csv.DictReader(open(journal_file, encoding="utf-8")))
            wins = sum(1 for r in rows if float(r.get("net_pnl", 0) or 0) > 0)
            result["trades"]   = len(rows)
            result["wins"]     = wins
            result["win_rate"] = round(wins / len(rows) * 100, 1) if rows else 0.0
        except Exception:
            pass
    comp_file = Path(BASE) / "comparison_log.json"
    if comp_file.exists():
        try:
            history = json.loads(comp_file.read_text(encoding="utf-8"))
            if history:
                last = history[-1]
                result["comparison_verdict"] = last.get("verdict", "")
                result["comparison_delta"]   = last.get("delta", {}).get("return_pct", 0.0)
                result["regime"]             = last.get("shadow", {}).get("regime", "UNKNOWN")
        except Exception:
            pass
    return result

# ── 3. METRICS ────────────────────────────────────────────────────────────────
def compute(logs, wb):
    today = date.today()
    if not logs:
        return None
    ops = logs[-1]

    regime  = ops.get("regime", "UNKNOWN")
    nifty   = ops.get("nifty", 0)
    dma50   = ops.get("dma50", 0)
    dma200  = ops.get("dma200", 0)

    # Days in current regime from consecutive tail of logs
    regime_days = 1
    for log in reversed(logs[:-1]):
        if log.get("regime") == regime:
            regime_days += 1
        else:
            break

    dist_pts = nifty - dma200
    dist_pct = (dist_pts / dma200 * 100) if dma200 else 0

    # Next rebalance
    future = sorted(d for d in wb["rebalance_dates"] if d > today)
    next_rebal    = future[0] if future else None
    days_to_rebal = (next_rebal - today).days if next_rebal else None

    # Paper day
    paper_day = max(1, (today - PAPER_START).days + 1) if today >= PAPER_START else 0

    # Portfolio
    port_vals = [r["portfolio"] for r in wb["daily_log"] if r["portfolio"] is not None]
    port_val  = port_vals[-1] if port_vals else PAPER_CAPITAL
    peak_val  = max(port_vals) if port_vals else PAPER_CAPITAL
    peak_val  = max(peak_val, PAPER_CAPITAL)

    # Cash
    alloc = sum(p["alloc"] for p in wb["positions"] if isinstance(p["alloc"], (int, float)))
    cash  = port_val - alloc

    # Drawdown
    drawdown = (peak_val - port_val) / peak_val * 100 if peak_val > 0 else 0

    # Unrealized
    unreal = sum(p["unrealized"] for p in wb["positions"] if isinstance(p["unrealized"], (int, float)))

    # Kill switches
    kills = []
    if regime == "BEAR" and wb["positions"]:
        kills.append(("K2", "Bear regime with open positions — exit immediately", "critical"))
    if drawdown >= 15:
        kills.append(("K3", f"Portfolio drawdown {drawdown:.1f}% >= 15% halt threshold", "critical"))
    elif drawdown >= 10:
        kills.append(("K1", f"Portfolio drawdown {drawdown:.1f}% >= 10% warning threshold", "warning"))

    # Completed trade stats
    completed = wb["trades"]
    wins = sum(1 for t in completed if t.get("win") == "Y")
    net_pnl = sum(t["net_pnl"] for t in completed if isinstance(t.get("net_pnl"), (int, float)))

    emergency = "CRITICAL" if any(k[2] == "critical" for k in kills) else \
                "WARNING"  if kills else "NOMINAL"

    shadow = load_shadow_data()

    return {
        "generated_at":  datetime.now().strftime("%Y-%m-%d %H:%M"),
        "data_date":     ops.get("data_date", str(today)),
        "paper_day":     paper_day,
        "regime":        regime,
        "regime_days":   regime_days,
        "nifty":         nifty,
        "dma50":         dma50,
        "dma200":        dma200,
        "dist_pts":      dist_pts,
        "dist_pct":      dist_pct,
        "next_rebal":    next_rebal,
        "days_to_rebal": days_to_rebal,
        "port_val":      port_val,
        "peak_val":      peak_val,
        "cash":          cash,
        "alloc":         alloc,
        "drawdown":      drawdown,
        "unreal":        unreal,
        "positions":     wb["positions"],
        "n_trades":      len(completed),
        "wins":          wins,
        "net_pnl":       net_pnl,
        "kills":         kills,
        "emergency":     emergency,
        "shadow":        shadow,
    }

# ── 4. HTML ───────────────────────────────────────────────────────────────────
def _inr(v):
    if v is None: return "—"
    try:
        v = float(v)
        sign = "-" if v < 0 else ("+" if v != 0 else "")
        return f"{sign}Rs.{abs(v):,.0f}"
    except Exception:
        return str(v)

def _d(d):
    if d is None: return "—"
    if isinstance(d, (date, datetime)): return d.strftime("%d %b %Y")
    return str(d)

REGIME_PALETTE = {
    "BEAR": {"accent": "#e05252", "bg": "#200e0e", "border": "#4a1a1a"},
    "BULL": {"accent": "#4caf50", "bg": "#0a1a0c", "border": "#1a3d1e"},
    "FLAT": {"accent": "#f0b429", "bg": "#1a1400", "border": "#3d3000"},
}

def render(m):
    r        = m["regime"]
    pal      = REGIME_PALETTE.get(r, {"accent":"#888","bg":"#1a1a1a","border":"#333"})
    em       = m["emergency"]
    em_color = {"NOMINAL":"#4caf50","WARNING":"#f0b429","CRITICAL":"#e05252"}[em]
    em_icon  = {"NOMINAL":"&#10003;","WARNING":"&#9888;","CRITICAL":"&#9888;"}[em]

    dist_dir = "above" if m["dist_pts"] >= 0 else "below"
    dist_c   = "#4caf50" if m["dist_pts"] >= 0 else "#e05252"
    dist_sgn = "+" if m["dist_pts"] >= 0 else ""

    dd_c = "#e05252" if m["drawdown"] >= 10 else ("#f0b429" if m["drawdown"] >= 5 else "#aaa")

    # Positions table rows
    pos_rows = ""
    for p in m["positions"]:
        unreal = p.get("unrealized", 0) or 0
        pct    = p.get("pnl_pct")
        pnl_c  = "#4caf50" if unreal >= 0 else "#e05252"
        cur    = p.get("current")
        sl     = p.get("sl")
        buf    = f"{(cur-sl)/cur*100:.1f}%" if sl and cur and cur > 0 else "—"
        pos_rows += (
            f"<tr>"
            f"<td class='td-sym'>{p['symbol']}</td>"
            f"<td class='td'>{_d(p.get('entry_date'))}</td>"
            f"<td class='td'>Rs.{float(p['entry_px']):,.2f}</td>"
            f"<td class='td'>{p.get('qty','—')}</td>"
            f"<td class='td'>{('Rs.'+f'{float(cur):,.2f}') if cur else '—'}</td>"
            f"<td class='td' style='color:{pnl_c};font-weight:500'>{_inr(unreal)}</td>"
            f"<td class='td' style='color:{pnl_c}'>{f'{float(pct)*100:.1f}%' if pct else '—'}</td>"
            f"<td class='td'>{('Rs.'+f'{float(sl):,.2f}') if sl else '—'} <span class='muted'>({buf} buffer)</span></td>"
            f"<td class='td muted'>{p.get('days_held','—')}</td>"
            f"</tr>"
        )
    if not pos_rows:
        msg = "No open positions — Bear regime. 100% cash." if r == "BEAR" else "No open positions."
        pos_rows = f"<tr><td colspan='9' class='td' style='text-align:center;color:#444;padding:20px'>{msg}</td></tr>"

    # Kill switch pills — K1 through K6
    kill_map = {k[0]: k for k in m["kills"]}
    pills = ""
    for k in ["K1","K2","K3","K4","K5","K6"]:
        if k in kill_map:
            kc = "#e05252" if kill_map[k][2]=="critical" else "#f0b429"
            pills += f"<span class='pill' style='border-color:{kc};color:{kc}' title='{kill_map[k][1]}'>{k} &#9888;</span>"
        else:
            pills += f"<span class='pill ok'>{k} &#10003;</span>"

    kill_detail = "".join(
        "<div class='kill-msg' style='color:{c}'>&#9888; {k}: {msg}</div>".format(
            c="#e05252" if sev == "critical" else "#f0b429", k=kid, msg=msg
        )
        for kid, msg, sev in m["kills"]
    )

    wr_str = f"{m['wins']}/{m['n_trades']} W/L" if m["n_trades"] > 0 else "No closed trades"

    sh = m["shadow"]
    sh_ret_c = "#4caf50" if sh["return_pct"] >= 0 else "#e05252"
    sh_pnl_c = "#4caf50" if sh["pnl"] >= 0 else "#e05252"
    sh_delta_c = "#4caf50" if sh["comparison_delta"] >= 0 else "#e05252"
    sh_regime_label = {
        "ACTIVE": "ACTIVE (50-DMA filter)",
        "CASH": "CASH (below 50-DMA)",
    }.get(sh["regime"], sh["regime"])

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta http-equiv="refresh" content="300">
<title>MoneyBot — Operator Dashboard</title>
<style>
*{{box-sizing:border-box;margin:0;padding:0}}
body{{font-family:'Segoe UI',system-ui,sans-serif;background:#0d0d0d;color:#ddd;padding:16px 20px;min-height:100vh}}
.topbar{{display:flex;justify-content:space-between;align-items:center;border-bottom:1px solid #1e1e1e;padding-bottom:12px;margin-bottom:16px}}
.title{{font-size:15px;font-weight:600;color:#fff;letter-spacing:.2px}}
.subtitle{{font-size:11px;color:#444;margin-top:3px}}
.em{{padding:5px 14px;border-radius:4px;font-size:12px;font-weight:600;letter-spacing:.6px;border:1px solid {em_color}44;color:{em_color};background:{em_color}11}}
.row3{{display:grid;grid-template-columns:repeat(3,1fr);gap:12px;margin-bottom:12px}}
.row4{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin-bottom:14px}}
.card{{background:#161616;border:1px solid #202020;border-radius:8px;padding:14px 16px}}
.clabel{{font-size:10px;color:#444;text-transform:uppercase;letter-spacing:.8px;margin-bottom:7px}}
.cval{{font-size:22px;font-weight:600;line-height:1.1;color:#e0e0e0}}
.csub{{font-size:12px;color:#444;margin-top:5px}}
.regime-card{{background:{pal['bg']};border:1px solid {pal['border']};border-radius:8px;padding:16px 18px}}
.regime-name{{font-size:38px;font-weight:700;color:{pal['accent']};letter-spacing:1px;line-height:1}}
.regime-desc{{font-size:12px;color:{pal['accent']}77;margin-top:8px}}
.regime-days{{font-size:11px;color:{pal['accent']}55;margin-top:6px}}
.section{{font-size:10px;color:#333;text-transform:uppercase;letter-spacing:.8px;margin:0 0 8px 2px}}
.pos-wrap{{background:#161616;border:1px solid #202020;border-radius:8px;overflow:hidden;margin-bottom:14px}}
table{{width:100%;border-collapse:collapse;font-size:13px}}
.th{{padding:7px 12px;text-align:left;color:#333;font-size:10px;text-transform:uppercase;letter-spacing:.6px;font-weight:500;background:#111;border-bottom:1px solid #1e1e1e}}
.td{{padding:9px 12px;border-bottom:1px solid #1a1a1a;color:#ccc;vertical-align:middle}}
.td-sym{{padding:9px 12px;border-bottom:1px solid #1a1a1a;font-weight:600;color:#e8e8e8}}
tbody tr:last-child td{{border-bottom:none}}
tbody tr:hover td, tbody tr:hover .td-sym{{background:#1a1a1a}}
.muted{{color:#444;font-size:11px}}
.kill-row{{background:#111;border:1px solid #1e1e1e;border-radius:8px;padding:10px 14px;display:flex;align-items:center;gap:8px;flex-wrap:wrap}}
.kill-label{{font-size:10px;color:#333;text-transform:uppercase;letter-spacing:.8px;margin-right:4px}}
.pill{{padding:3px 10px;border-radius:4px;border:1px solid #2a2a2a;font-size:11px;font-weight:500;white-space:nowrap}}
.pill.ok{{color:#1e4d24;border-color:#1a2e1c}}
.kill-msg{{font-size:12px;margin-top:5px;padding-left:2px}}
</style>
</head>
<body>

<div class="topbar">
  <div>
    <div class="title">MoneyBot RS60/EP40 &nbsp;&#183;&nbsp; Operator Dashboard</div>
    <div class="subtitle">Paper day {m['paper_day']} &nbsp;&#183;&nbsp; Data: {m['data_date']} &nbsp;&#183;&nbsp; Generated: {m['generated_at']} &nbsp;&#183;&nbsp; Auto-refresh 5 min</div>
  </div>
  <div class="em">{em_icon} &nbsp;{em}</div>
</div>

<div class="row3">

  <div class="regime-card">
    <div class="clabel" style="color:{pal['accent']}55">Regime &nbsp;(Policy C)</div>
    <div class="regime-name">{r}</div>
    <div class="regime-desc">{'100% cash &mdash; no new entries permitted' if r=='BEAR' else 'Tradeable &mdash; run candidate list on rebalance day'}</div>
    <div class="regime-days">{m['regime_days']} consecutive day{'s' if m['regime_days']!=1 else ''} tracked in this regime</div>
  </div>

  <div class="card">
    <div class="clabel">Distance to 200-DMA</div>
    <div class="cval" style="color:{dist_c};font-size:28px">{dist_sgn}{abs(m['dist_pts']):,.0f} <span style="font-size:15px;font-weight:400;color:{dist_c}88">pts</span></div>
    <div class="csub">Nifty is <span style="color:{dist_c}">{abs(m['dist_pct']):.2f}%</span> {dist_dir} the 200-DMA</div>
    <div style="margin-top:10px;font-size:11px;color:#2a2a2a;display:flex;gap:16px">
      <span>Nifty <span style="color:#555">{m['nifty']:,.0f}</span></span>
      <span>50-DMA <span style="color:#555">{m['dma50']:,.0f}</span></span>
      <span>200-DMA <span style="color:#555">{m['dma200']:,.0f}</span></span>
    </div>
  </div>

  <div class="card">
    <div class="clabel">Next rebalance</div>
    <div class="cval" style="font-size:20px">{_d(m['next_rebal'])}</div>
    <div class="csub">{'in ' + str(m['days_to_rebal']) + ' calendar days' if m['days_to_rebal'] else 'None scheduled'}</div>
    <div style="margin-top:10px;font-size:11px;{'color:#f0b429' if m['days_to_rebal']==0 else 'color:#2a2a2a'}">
      {'&#9654; Today is rebalance day &mdash; check Top-10 before 9:20 AM' if m['days_to_rebal']==0 else f'Completed trades: {m["n_trades"]} &nbsp;&#183;&nbsp; {wr_str}'}
    </div>
  </div>

</div>

<div class="row4">

  <div class="card">
    <div class="clabel">Portfolio value</div>
    <div class="cval">Rs.{m['port_val']:,.0f}</div>
    <div class="csub">Started Rs.{PAPER_CAPITAL:,}</div>
  </div>

  <div class="card">
    <div class="clabel">Cash available</div>
    <div class="cval">Rs.{m['cash']:,.0f}</div>
    <div class="csub">{f"Rs.{m['alloc']:,.0f} in {len(m['positions'])} position{'s' if len(m['positions'])!=1 else ''}" if m['positions'] else "No open positions"}</div>
  </div>

  <div class="card">
    <div class="clabel">Drawdown from peak</div>
    <div class="cval" style="color:{dd_c}">{'-' if m['drawdown']>0 else ''}{m['drawdown']:.1f}%</div>
    <div class="csub">Peak Rs.{m['peak_val']:,.0f}</div>
  </div>

  <div class="card">
    <div class="clabel">Unrealized P&amp;L</div>
    <div class="cval" style="color:{'#4caf50' if m['unreal']>=0 else '#e05252'}">{'+' if m['unreal']>0 else ''}Rs.{m['unreal']:,.0f}</div>
    <div class="csub">{len(m['positions'])} open position{'s' if len(m['positions'])!=1 else ''}</div>
  </div>

</div>

<div class="section">Open positions ({len(m['positions'])})</div>
<div class="pos-wrap">
  <table>
    <thead>
      <tr>
        <th class="th">Symbol</th>
        <th class="th">Entry date</th>
        <th class="th">Entry price</th>
        <th class="th">Qty</th>
        <th class="th">Current price</th>
        <th class="th">Unrealized P&amp;L</th>
        <th class="th">P&amp;L %</th>
        <th class="th">Stop-loss</th>
        <th class="th">Days held</th>
      </tr>
    </thead>
    <tbody>{pos_rows}</tbody>
  </table>
</div>

<div class="section">Shadow Engine &mdash; Experimental (RS60/EP40 + 50-DMA regime, paper only)</div>
<div style="background:#111;border:1px solid #1a1a2a;border-radius:8px;padding:14px 16px;margin-bottom:14px;display:grid;grid-template-columns:repeat(5,1fr);gap:12px">
  <div>
    <div class="clabel" style="color:#5566aa">Regime</div>
    <div style="font-size:15px;font-weight:600;color:#7788cc">{sh_regime_label}</div>
  </div>
  <div>
    <div class="clabel" style="color:#5566aa">Portfolio value</div>
    <div style="font-size:15px;font-weight:600;color:#aabbdd">Rs.{sh['portfolio_value']:,.0f}</div>
  </div>
  <div>
    <div class="clabel" style="color:#5566aa">Total PnL</div>
    <div style="font-size:15px;font-weight:600;color:{sh_pnl_c}">{'+' if sh['pnl']>=0 else ''}Rs.{sh['pnl']:,.0f}</div>
  </div>
  <div>
    <div class="clabel" style="color:#5566aa">Return</div>
    <div style="font-size:15px;font-weight:600;color:{sh_ret_c}">{sh['return_pct']:+.2f}%</div>
    <div style="font-size:10px;color:#334;margin-top:3px">vs prod {sh['comparison_delta']:+.2f}% <span style="color:{sh_delta_c}">delta</span></div>
  </div>
  <div>
    <div class="clabel" style="color:#5566aa">Trades / Win rate</div>
    <div style="font-size:15px;font-weight:600;color:#aabbdd">{sh['trades']} trades</div>
    <div style="font-size:11px;color:#5566aa;margin-top:3px">{sh['win_rate']:.1f}% win rate</div>
  </div>
</div>
{f'<div style="font-size:11px;color:#445;margin:-8px 0 12px 2px;padding:6px 10px;background:#0e0e14;border-radius:4px;border-left:2px solid #334">&nbsp;{sh["comparison_verdict"]}</div>' if sh["comparison_verdict"] else ''}

<div class="kill-row">
  <span class="kill-label">Kill switches</span>
  {pills}
  {'<span class="muted" style="margin-left:6px;font-size:11px">All checks passed</span>' if not m['kills'] else ''}
</div>
{kill_detail}

</body>
</html>"""

# ── 5. MAIN ───────────────────────────────────────────────────────────────────

def _build_and_write():
    """Shared logic: load data, compute metrics, write HTML. Returns output path."""
    logs = load_ops_logs()
    if not logs:
        print(f"ERROR: No ops_log_*.json files found in {BASE}")
        return None

    wb = load_workbook()
    m  = compute(logs, wb)
    if m is None:
        print("ERROR: Could not compute metrics — no valid ops logs.")
        return None

    print(f"  Regime: {m['regime']}  |  Nifty: {m['nifty']:,.0f}  |  Emergency: {m['emergency']}")
    print(f"  Positions: {len(m['positions'])}  |  Cash: Rs.{m['cash']:,.0f}  |  DD: {m['drawdown']:.1f}%")

    html     = render(m)
    out_path = os.path.join(BASE, "dashboard.html")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)
    return out_path


def main():
    """Build dashboard and open in browser (first call only)."""
    print("MoneyBot Dashboard — loading...")
    out_path = _build_and_write()
    if out_path is None:
        sys.exit(1)
    webbrowser.open(f"file:///{out_path.replace(os.sep, '/')}")
    print(f"  Opened: {out_path}")


def refresh_only():
    """Rebuild dashboard HTML without opening a new browser tab.
    The HTML has <meta http-equiv='refresh' content='300'> so the
    already-open tab will pick up the new file automatically."""
    out_path = _build_and_write()
    if out_path:
        print(f"  Dashboard refreshed: {out_path}")


if __name__ == "__main__":
    main()
